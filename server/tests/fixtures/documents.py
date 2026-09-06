"""Generated fixtures for every supported format.

Generated rather than committed, for the same reason as the PDF fixtures: the expected
content lives in code beside the assertions. Anubhav's real sample documents
go in ``samples/`` and are exercised by the seed route, not by these tests, so the unit
suite stays fast and deterministic.
"""

from __future__ import annotations

import io

from docx import Document as DocxDocument
from openpyxl import Workbook
from PIL import Image, ImageDraw


def docx_invoice() -> bytes:
    """A Word document shaped like an invoice, with a heading and paragraphs."""
    buffer = io.BytesIO()
    document = DocxDocument()
    document.add_heading("Invoice", level=1)
    document.add_paragraph("Supplier: Contoso Supplies Limited")
    document.add_paragraph("Invoice Number: CS-4471")
    document.add_paragraph("Issue Date: 2026-02-11")
    document.add_paragraph("Total Due: EUR 3,410.00")
    document.add_paragraph(
        "Notes: this covers the February retainer and two ad hoc support incidents "
        "raised on the eighth and the ninth."
    )
    document.save(buffer)
    return buffer.getvalue()


def xlsx_ledger() -> bytes:
    """A spreadsheet with two sheets, so sheet-level provenance can be tested."""
    buffer = io.BytesIO()
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Invoices"
    first.append(["Vendor", "Invoice No", "Amount", "Currency", "Date"])
    first.append(["Acme Industrial", "AI-9001", 1250.75, "USD", "2026-01-15"])
    first.append(["Globex Freight", "GF-2231", 480.00, "USD", "2026-01-22"])

    second = workbook.create_sheet("Notes")
    second.append(["Reference", "Comment"])
    second.append(["AI-9001", "Awaiting purchase order confirmation"])

    workbook.save(buffer)
    return buffer.getvalue()


CSV_STATEMENT = (
    b"Date,Description,Debit,Credit,Balance\n"
    b"2026-01-03,Opening balance,,,10000.00\n"
    b"2026-01-11,ACME INDUSTRIAL INV AI-9001,1250.75,,8749.25\n"
    b"2026-01-19,Client receipt,,4000.00,12749.25\n"
)
"""A bank statement, which deliberately has NO vendor field. This is the document that
makes the unification problem real: it cannot be forced into an invoice schema, so it is
what schema drift detection has to handle gracefully."""

TXT_RECEIPT = (
    b"CORNER CAFE\n"
    b"12 Mill Road\n"
    b"\n"
    b"Receipt 00841\n"
    b"Date: 04 Feb 2026\n"
    b"\n"
    b"2 x Flat white        7.00\n"
    b"1 x Sandwich          6.50\n"
    b"\n"
    b"Total                13.50 GBP\n"
    b"Paid by card\n"
)

PROSE_TXT = (
    b"This is a paragraph of ordinary prose, with commas in it, several of them, "
    b"which must not be mistaken for a spreadsheet.\n"
    b"It runs across more than one line, and the lines are of different lengths.\n"
)


def png_receipt(width: int = 600, height: int = 320) -> bytes:
    """An image with rendered text, for the text-recognition path."""
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    lines = [
        "HARDWARE DEPOT",
        "Invoice: HD-7788",
        "Date: 2026-03-02",
        "Total: $842.10",
    ]
    for row, line in enumerate(lines):
        draw.text((20, 30 + row * 40), line, fill="black")
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()
