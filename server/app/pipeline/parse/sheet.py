"""Spreadsheet and comma-separated value parsing.

One sheet becomes one or more pages, and every row becomes a line whose ``locator`` is its
address, such as ``Invoices!row 3``. A model quoting a single cell's value still gets a
highlight box around just that cell's text, because grounding matches the quote against the
word boxes within the row rather than against the row as a unit.

Empty trailing rows and columns are trimmed: spreadsheets routinely report a used range far
larger than their content, and rendering three hundred blank rows would bury the data in
whitespace and waste pages.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

from app.config import Settings
from app.domain.document import ParsedDocument, ParsedPage, SourceFormat
from app.errors import ParseFailed
from app.logging import get_logger
from app.pipeline.parse.pdf import ProgressCallback
from app.pipeline.parse.render import RenderLine, render_pages

logger = get_logger(__name__)

MAX_ROWS_PER_SHEET = 2000
MAX_COLUMNS = 40


def _format_cell(value: Any) -> str:
    """Render a cell for the page, preserving what a person would see.

    Dates are written as ISO strings rather than as Python repr output, and whole-number
    floats lose their trailing ``.0``, because the rendered page is what the model reads
    AND what the user sees highlighted. A value that reads differently in the two places
    would make provenance look wrong even when it is right.
    """
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else (
            value.isoformat(sep=" ", timespec="minutes")
        )
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _trim(rows: list[list[str]]) -> list[list[str]]:
    """Drop trailing empty rows and columns."""
    while rows and not any(cell for cell in rows[-1]):
        rows.pop()
    if not rows:
        return []
    width = max(
        (max((index + 1 for index, cell in enumerate(row) if cell), default=0) for row in rows),
        default=0,
    )
    return [row[:width] for row in rows]


def _sheet_lines(name: str, rows: list[list[str]]) -> list[RenderLine]:
    return [
        RenderLine(" | ".join(row), f"{name}!row {number}")
        for number, row in enumerate(rows, start=1)
        if any(cell for cell in row)
    ]


def _pages_from_sheets(
    sheets: list[tuple[str, list[list[str]]]], dpi: int
) -> tuple[list[ParsedPage], list[bytes]]:
    """Render each sheet separately, then renumber pages across the whole document."""
    pages: list[ParsedPage] = []
    images: list[bytes] = []

    for name, rows in sheets:
        lines = _sheet_lines(name, rows)
        if not lines:
            continue
        for page, png in render_pages(lines, dpi=dpi):
            # `render_pages` numbers from zero per call, so reindex and record which sheet
            # this page came from.
            pages.append(page.model_copy(update={"index": len(pages), "locator": name}))
            images.append(png)

    return pages, images


def parse_xlsx(
    data: bytes, settings: Settings, on_progress: ProgressCallback | None = None
) -> tuple[ParsedDocument, list[bytes]]:
    report = on_progress or (lambda _detail: None)
    report("reading spreadsheet")

    try:
        # read_only streams rows instead of building the whole grid in memory;
        # data_only returns computed values rather than formula text, which is what a
        # reader of the document sees.
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseFailed(
            "We could not read this spreadsheet. It may be corrupt, or saved in the older "
            ".xls format, which Distill does not support.",
            library_error=type(exc).__name__,
        ) from exc

    sheets: list[tuple[str, list[list[str]]]] = []
    try:
        for worksheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if row_index >= MAX_ROWS_PER_SHEET:
                    report(
                        f"sheet {worksheet.title!r} has more than "
                        f"{MAX_ROWS_PER_SHEET} rows, reading the first {MAX_ROWS_PER_SHEET}"
                    )
                    break
                rows.append([_format_cell(cell) for cell in row[:MAX_COLUMNS]])
            trimmed = _trim(rows)
            if trimmed:
                sheets.append((worksheet.title, trimmed))
    finally:
        workbook.close()

    if not sheets:
        raise ParseFailed("This spreadsheet has no data in any of its sheets.")

    report("laying out pages")
    pages, images = _pages_from_sheets(sheets, settings.page_render_dpi)
    parsed = ParsedDocument(source_format=SourceFormat.XLSX, pages=pages)
    logger.info("xlsx.parsed", sheets=len(sheets), pages=parsed.page_count)
    return parsed, images


def parse_csv(
    data: bytes, settings: Settings, on_progress: ProgressCallback | None = None
) -> tuple[ParsedDocument, list[bytes]]:
    report = on_progress or (lambda _detail: None)
    report("reading tabular data")

    text: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    if text is None:
        raise ParseFailed("We could not read the text encoding of this file.")

    sample = "\n".join(text.splitlines()[:20])
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        # Falling back rather than failing: a single-column file has no delimiter to find,
        # and is still perfectly readable as one column per row.
        dialect = csv.excel

    rows = [
        [str(cell).strip() for cell in row[:MAX_COLUMNS]]
        for row in csv.reader(io.StringIO(text), dialect)
    ][:MAX_ROWS_PER_SHEET]
    trimmed = _trim(rows)
    if not trimmed:
        raise ParseFailed("This file has no rows of data in it.")

    report("laying out pages")
    pages, images = _pages_from_sheets([("Data", trimmed)], settings.page_render_dpi)
    parsed = ParsedDocument(source_format=SourceFormat.CSV, pages=pages)
    logger.info("csv.parsed", rows=len(trimmed), pages=parsed.page_count)
    return parsed, images
